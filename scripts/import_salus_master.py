#!/usr/bin/env python3
"""
Batch 248 — SALUS master-data importer (meal-plan customer).

Salus_Master_Recipes_with_Inventory_Codes.xlsx has the same "Raw material list"
and "Recipe Ingredients" layout as FRSH / SMC, plus three things they don't:

  1. PLANS. Menu sheet, packages table beside the menu (Packages / protein /
     Carb): COMFY 120g/150g, Low Carb 120g/100g, Grit 150g/200g, Salus Fit
     80/100g.                                              -> customer_plans
  2. WHICH DISHES ARE IN WHICH PLAN. Menu sheet, the four cells under "Plan"
     (Comfy / Grit / Low Carb / Salus Fit) on each main dish.  -> recipe_plans
  3. PER-PLAN QUANTITIES. Recipe Ingredients, columns "Sulus Fit", "Low Carb",
     "COMFY", "GRIT" on the plan-sensitive lines (the protein), per batch —
     e.g. Butter Chicken chicken breast 1300 / 1800 / 1800 / 2300 g per 10.
                                                    -> recipe_ingredient_plans

WORKBOOK NOTE (read this): on those plan lines the workbook's base "NET Qty
req per Batch" is the SUM of the four plan quantities (Butter Chicken 7200 =
1300+1800+1800+2300), so the recipe's base food cost is costed at ~4x the
protein a real plate gets. Orders placed for a plan are costed, BOM'd and
issued at the plan quantity, so this does not reach production. The base
figure only shows on the recipe master itself; the importer prints the list.

Raw materials and recipe lines go through the SAME functions as the FRSH/SMC
importer (scripts/import_frsh_master.py), so column mapping, section mapping
and the idempotent upsert behave identically.

USAGE
    python scripts/import_salus_master.py --file "Salus_Master_Recipes_with_Inventory_Codes.xlsx" --company 1 --dry-run
    python scripts/import_salus_master.py --file "Salus_Master_Recipes_with_Inventory_Codes.xlsx" --company 1 \
        --brand "Gourmet 360" --channel "Corporate"

--brand / --channel (optional) set the Salus customer's default brand and
sales channel, so the order form fills them the moment Salus is picked.
"""
from __future__ import annotations

import argparse
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)

import openpyxl  # noqa: E402
from sqlalchemy import text  # noqa: E402

import import_frsh_master as base  # noqa: E402  (shared column/section logic)

try:
    from app.database.session import SessionLocal
except Exception:  # pragma: no cover
    from app.db import SessionLocal  # type: ignore

CUSTOMER_KEY = "SALUS"
DAY_ORDER = ["Saturday", "Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

# Canonical plan codes. The workbook spells each plan three ways
# ("Sulus Fit" / "Salus Fit", "Low Carb ( L C)" / "Low Carb", "COMFY" / "Comfy").
PLANS = [
    # code,       display,     sort, match stems (letters only, lowercase)
    ("COMFY",     "Comfy",     1, ("comfy",)),
    ("LOW_CARB",  "Low Carb",  2, ("lowcarb",)),
    ("GRIT",      "Grit",      3, ("grit",)),
    ("SALUS_FIT", "Salus Fit", 4, ("salusfit", "sulusfit", "fit")),
]


def plan_code(raw) -> str:
    key = re.sub(r"[^a-z]", "", str(raw or "").lower())
    if not key:
        return ""
    for code, _name, _sort, stems in PLANS:
        if any(key.startswith(s) for s in stems):
            return code
    return ""


def _grams(v) -> float | None:
    m = re.search(r"[\d.]+", str(v or ""))
    return float(m.group(0)) if m else None


# ---------------------------------------------------------------------------
# 1. Plans (packages table on the Menu sheet)
# ---------------------------------------------------------------------------
def read_packages(ws) -> dict[str, dict]:
    """Find the 'Packages | protein | Carb' block anywhere on the sheet."""
    out: dict[str, dict] = {}
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 60, 60), values_only=True))
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            if base._s(v).lower() == "packages":
                for r2 in rows[ri + 1: ri + 12]:
                    if ci >= len(r2) or not base._s(r2[ci]):
                        break
                    code = plan_code(r2[ci])
                    if code:
                        out[code] = {"protein": _grams(r2[ci + 1] if ci + 1 < len(r2) else None),
                                     "carb": _grams(r2[ci + 2] if ci + 2 < len(r2) else None)}
                return out
    return out


def upsert_plans(db, company_id, packages, dry) -> int:
    n = 0
    for code, name, sort, _ in PLANS:
        pk = packages.get(code, {})
        n += 1
        if dry:
            continue
        db.execute(text("""
            INSERT INTO customer_plans
                (company_id, customer_key, plan_code, plan_name, protein_g, carb_g,
                 sort_order, is_active, created_at, updated_at)
            VALUES (:cid, :key, :code, :name, :p, :c, :sort, 1, NOW(), NOW())
            ON DUPLICATE KEY UPDATE plan_name = VALUES(plan_name),
                protein_g = COALESCE(VALUES(protein_g), customer_plans.protein_g),
                carb_g = COALESCE(VALUES(carb_g), customer_plans.carb_g),
                sort_order = VALUES(sort_order), is_active = 1, updated_at = NOW()
        """), {"cid": company_id, "key": CUSTOMER_KEY, "code": code, "name": name,
               "p": pk.get("protein"), "c": pk.get("carb"), "sort": sort})
    if not dry:
        db.commit()
    return n


# ---------------------------------------------------------------------------
# 2. Menu: days + category + which plans each dish is in
# ---------------------------------------------------------------------------
def import_menu(db, ws, company_id, dry) -> dict:
    hrow, hdr = base._find_header(ws, ["recipe code", "days"])
    if hrow is None:
        print("  ! Menu header not found — skipping")
        return {"recipes": 0, "no_code": 0, "plan_links": 0}
    c_day, c_code = base.col(hdr, "days"), base.col(hdr, "recipe code")
    c_name = base.col(hdr, "name (en)", "recipe names", "recipe name")
    c_cat = base.col(hdr, "category")
    c_plan = base.col(hdr, "plan")

    agg: dict[str, dict] = {}
    no_code: list[str] = []
    for row in ws.iter_rows(min_row=hrow + 1, values_only=True):
        code = base._s(row[c_code]) if c_code is not None else ""
        name = base._s(row[c_name]) if c_name is not None else ""
        if not code:
            if name:
                no_code.append(name)
            continue
        a = agg.setdefault(code, {"name": name, "cat": "", "days": set(), "plans": set()})
        day = base._s(row[c_day]) if c_day is not None else ""
        if day:
            a["days"].add(day.title())
        cat = base._s(row[c_cat]) if c_cat is not None else ""
        if cat and not a["cat"]:
            a["cat"] = cat
        # "Plan" heads a block of four cells (Comfy / Grit / Low Carb / Salus Fit).
        if c_plan is not None:
            for v in row[c_plan:c_plan + 4]:
                pc = plan_code(v)
                if pc:
                    a["plans"].add(pc)

    links = 0
    if not dry:
        # Same reason as the FRSH importer: clear stale days first, so what is
        # stored is exactly this workbook's menu.
        db.execute(text("""
            UPDATE recipes SET day_of_week = NULL, updated_at = NOW()
            WHERE company_id = :cid AND UPPER(customer_name) = :key
        """), {"cid": company_id, "key": CUSTOMER_KEY})
        db.execute(text("DELETE FROM recipe_plans WHERE company_id = :cid AND recipe_code IN "
                        "(SELECT recipe_code FROM recipes WHERE company_id = :cid "
                        " AND UPPER(customer_name) = :key)"),
                   {"cid": company_id, "key": CUSTOMER_KEY})
    for code, a in agg.items():
        links += len(a["plans"])
        if dry:
            continue
        days = " & ".join(d for d in DAY_ORDER if d in a["days"])
        db.execute(text("""
            UPDATE recipes SET day_of_week = :day,
                   category = COALESCE(NULLIF(category,''), :cat), updated_at = NOW()
            WHERE company_id = :cid AND recipe_code = :code
        """), {"cid": company_id, "code": code, "day": days, "cat": a["cat"]})
        for pc in sorted(a["plans"]):
            db.execute(text("""
                INSERT IGNORE INTO recipe_plans (company_id, recipe_code, plan_code)
                VALUES (:cid, :code, :pc)
            """), {"cid": company_id, "code": code, "pc": pc})
    if not dry:
        db.commit()
    return {"recipes": len(agg), "no_code": len(no_code), "no_code_names": no_code,
            "plan_links": links}


# ---------------------------------------------------------------------------
# 3. Per-plan quantities on recipe lines
# ---------------------------------------------------------------------------
def import_plan_quantities(db, ws, company_id, dry) -> dict:
    """Second pass over Recipe Ingredients, after base.import_recipe_ingredients.

    That pass inserts one recipe_ingredients row per sheet row that has both a
    recipe ref and an item code, numbering them line_no 1..n in sheet order.
    This pass walks the sheet with the SAME filter and the same counter, so
    (recipe_code, n) identifies the exact row it created.
    """
    hrow, hdr = base._find_header(ws, [("recipe code", "recipe ref"), "item code"])
    c_rcode = base.col(hdr, "recipe code", "recipe ref")
    c_icode = base.col(hdr, "item code")
    plan_cols = {}
    for h, j in hdr.items():
        pc = plan_code(h)
        if pc and h not in ("plan",):
            plan_cols[pc] = j

    counters: dict[str, int] = {}
    wanted: list[tuple] = []           # (recipe_code, line_no, plan_code, qty)
    skipped_no_code = 0
    base_is_sum: list[str] = []
    for row in ws.iter_rows(min_row=hrow + 1, values_only=True):
        rcode = base._s(row[c_rcode]) if c_rcode is not None else ""
        icode = base._s(row[c_icode]) if c_icode is not None else ""
        vals = {pc: base._f(row[j], None) for pc, j in plan_cols.items()}
        vals = {pc: v for pc, v in vals.items() if v is not None and v > 0}
        if not rcode or not icode:
            if rcode and vals:
                skipped_no_code += 1
            continue
        counters[rcode] = counters.get(rcode, 0) + 1
        for pc, q in vals.items():
            wanted.append((rcode, counters[rcode], pc, q))
        if len(vals) > 1:
            c_net = base.col(hdr, "net qty req per batch (g/pcs)")
            net = base._f(row[c_net]) if c_net is not None else 0
            if net and abs(net - sum(vals.values())) < 0.01:
                base_is_sum.append(rcode)

    if not dry:
        # Orphans from the lines the base pass just deleted and re-inserted.
        db.execute(text("""
            DELETE rip FROM recipe_ingredient_plans rip
            LEFT JOIN recipe_ingredients ri ON ri.id = rip.recipe_ingredient_id
            WHERE ri.id IS NULL
        """))
        for rcode, ln, pc, q in wanted:
            db.execute(text("""
                INSERT INTO recipe_ingredient_plans (recipe_ingredient_id, plan_code, qty_batch)
                SELECT ri.id, :pc, :q
                FROM recipe_ingredients ri JOIN recipes r ON r.id = ri.recipe_id
                WHERE r.company_id = :cid AND r.recipe_code = :code AND ri.line_no = :ln
                ON DUPLICATE KEY UPDATE qty_batch = VALUES(qty_batch)
            """), {"pc": pc, "q": q, "cid": company_id, "code": rcode, "ln": ln})
        db.commit()
    return {"quantities": len(wanted), "lines": len({(r, l) for r, l, _, _ in wanted}),
            "skipped_no_item_code": skipped_no_code, "base_is_sum": sorted(set(base_is_sum))}


def set_customer_defaults(db, company_id, brand, channel) -> str:
    """Default brand/channel on every customer whose name contains 'Salus'."""
    def _resolve(table, code_col, name_col, v):
        if not v:
            return None
        r = db.execute(text(f"""
            SELECT {code_col} FROM {table}
            WHERE ({code_col} = :v OR {name_col} = :v) AND (company_id = :cid OR company_id IS NULL)
            LIMIT 1"""), {"v": v, "cid": company_id}).scalar()
        return r or None

    b = _resolve("brands", "brand_code", "brand_name_en", brand)
    c = _resolve("revenue_streams", "stream_code", "stream_name", channel)
    from app.modules.orders.routes_menu import ensure_default_columns
    ensure_default_columns(db)
    sets, params = [], {"cid": company_id}
    if b:
        sets.append("default_brand_code = :b"); params["b"] = b
    if c:
        sets.append("default_channel_code = :c"); params["c"] = c
    if not sets:
        return f"brand '{brand}' / channel '{channel}' not found in master — defaults not set"
    n = db.execute(text(f"""
        UPDATE customers SET {', '.join(sets)}
        WHERE (company_id = :cid OR company_id IS NULL) AND LOWER(customer_name) LIKE '%salus%'
    """), params).rowcount
    db.commit()
    return f"defaults set on {n} Salus customer(s): brand={b or '-'} channel={c or '-'}"


def recalc(db, company_id) -> int:
    from app.models.recipe import Recipe
    from app.services.recipe_service import recalc_recipe
    n = 0
    for r in db.query(Recipe).filter(Recipe.company_id == company_id,
                                     Recipe.customer_name.in_(["SALUS", "Salus", "salus"])).all():
        recalc_recipe(r)
        n += 1
    db.commit()
    return n


def main():
    ap = argparse.ArgumentParser(description="Import SALUS master data (with meal plans)")
    ap.add_argument("--file", required=True)
    ap.add_argument("--company", type=int, default=1)
    ap.add_argument("--brand", default="", help="Default brand for Salus (code or name)")
    ap.add_argument("--channel", default="", help="Default sales channel for Salus (code or name)")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    if not os.path.exists(args.file):
        print(f"File not found: {args.file}")
        sys.exit(1)

    wb = openpyxl.load_workbook(args.file, read_only=True, data_only=True)
    names = {n.lower().strip(): n for n in wb.sheetnames}
    db = SessionLocal()
    dry = args.dry_run
    try:
        print(f"{'DRY RUN — ' if dry else ''}Importing SALUS into company_id={args.company}\n")
        if not dry:
            from app.services.customer_plans import ensure_plan_schema
            ensure_plan_schema(db)

        if "raw material list" in names:
            ok, fail = base.import_raw_materials(db, wb[names["raw material list"]], args.company, dry)
            print(f"  Ingredients (raw materials): {ok} ok, {fail} failed")

        if "recipe ingredients" in names:
            ws = wb[names["recipe ingredients"]]
            rok, lok, rfail = base.import_recipe_ingredients(db, ws, args.company, dry)
            print(f"  Recipes: {rok} ok, {rfail} failed  |  Recipe lines: {lok}")
            pq = import_plan_quantities(db, ws, args.company, dry)
            print(f"  Plan quantities: {pq['quantities']} on {pq['lines']} line(s)")
            if pq["skipped_no_item_code"]:
                print(f"    ! {pq['skipped_no_item_code']} plan line(s) have no Item Code and were "
                      f"skipped (sub-recipe items) — give them a code to include them")
            if pq["base_is_sum"]:
                print(f"    ! {len(pq['base_is_sum'])} recipe(s) where the base NET qty = sum of the "
                      f"4 plans (base recipe cost overstated; plan orders unaffected):")
                print("      " + ", ".join(pq["base_is_sum"]))

        if "menu" in names:
            ws = wb[names["menu"]]
            pk = read_packages(ws)
            n = upsert_plans(db, args.company, pk, dry)
            desc = ", ".join(f"{c} protein {v.get('protein')}g / carb {v.get('carb')}g"
                             for c, v in pk.items())
            print(f"  Plans: {n} ({desc or 'packages table not found — names only'})")
            m = import_menu(db, ws, args.company, dry)
            print(f"  Menu: {m['recipes']} recipe(s), {m['plan_links']} recipe-plan link(s)")
            if m["no_code"]:
                print(f"    ! {m['no_code']} menu row(s) without a Recipe Code were skipped "
                      f"(e.g. {', '.join(sorted(set(m['no_code_names']))[:6])})")

        if not dry:
            print(f"  Recipe header costs recalculated: {recalc(db, args.company)}")
            if args.brand or args.channel:
                print("  " + set_customer_defaults(db, args.company, args.brand, args.channel))

        print("\nDry run complete — nothing written." if dry else "\nDone. SALUS master data imported.")
    finally:
        db.close()
        wb.close()


if __name__ == "__main__":
    main()
